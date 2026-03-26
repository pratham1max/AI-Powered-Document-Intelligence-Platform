"""
Celery tasks for the 'default' queue — parse, chunk, and embed documents.
"""
import asyncio
import logging
from pathlib import Path

from celery_app import app
from utils.crypto import decrypt_file
from utils.chunker import chunk_text
from utils.embeddings import get_embedding
from models.database import AsyncSessionLocal
from models.models import Document, DocumentChunk, Tag

import PyPDF2
import docx
import openpyxl
import io

logger = logging.getLogger(__name__)


def extract_text(file_bytes: bytes, mime_type: str) -> str:
    """Extract plain text from supported file types."""
    if mime_type == "application/pdf":
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    elif mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",):
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)
    elif mime_type in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                lines.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(lines)
    else:
        # Plain text fallback
        return file_bytes.decode("utf-8", errors="replace")


async def _process(document_id: int, user_id: int, enc_path: str, mime_type: str):
    async with AsyncSessionLocal() as session:
        doc = await session.get(Document, document_id)
        if not doc:
            return

        try:
            doc.status = "processing"
            await session.commit()

            file_bytes = decrypt_file(enc_path, user_id)
            text = extract_text(file_bytes, mime_type)
            chunks = chunk_text(text)

            for i, chunk_content in enumerate(chunks):
                embedding = get_embedding(chunk_content)
                chunk = DocumentChunk(
                    document_id=document_id,
                    chunk_index=i,
                    content=chunk_content,
                    embedding=embedding,  # list[float] — works for both JSON (SQLite) and Vector (pgvector)
                )
                session.add(chunk)

            doc.status = "ready"
            await session.commit()
            logger.info("Document %s processed: %d chunks", document_id, len(chunks))

        except Exception as exc:
            doc.status = "error"
            await session.commit()
            logger.error("Failed to process document %s: %s", document_id, exc)
            raise


@app.task(bind=True, queue="default", max_retries=2)
def task_parse_and_embed(self, document_id: int, user_id: int, enc_path: str, mime_type: str):
    """Parse, chunk, and embed a document after encryption."""
    try:
        asyncio.run(_process(document_id, user_id, enc_path, mime_type))
    except Exception as exc:
        raise self.retry(exc=exc)
